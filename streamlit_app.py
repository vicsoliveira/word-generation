import streamlit as st
import pandas as pd
from docx import Document
from io import BytesIO

def create_word_document(df):
    """
    Creates a Word document from a DataFrame.

    Parameters:
    df (pd.DataFrame): DataFrame containing the data to be added to the Word document.

    Returns:
    doc (Document): A Word document with the data from the DataFrame.
    """
    doc = Document()
    doc.add_heading('Report Generated from Excel', level=1)

    for i, row in df.iterrows():
        doc.add_heading(f'Row {i + 1}', level=2)
        for col in df.columns:
            doc.add_paragraph(f'{col}: {row[col]}')

    return doc

# Streamlit app layout
st.title('Excel to Word Document Generator')

# File uploader to upload Excel file
uploaded_file = st.file_uploader("Choose an Excel file", type="xlsx")

if uploaded_file is not None:
    # Read the Excel file into a DataFrame
    df = pd.read_excel(uploaded_file, engine='openpyxl')

    # Create a Word document from the DataFrame
    doc = create_word_document(df)

    # Save the document to a BytesIO object
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    st.success("Word document created successfully!")

    # Download button to download the Word document
    st.download_button(
        label="Download Word Document",
        data=buffer,
        file_name="report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
